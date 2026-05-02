"""
Day 28 - Practice Input File Generator
Creates: day28_practice_data.xlsx
Data: Monthly financial data for charting practice.

"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()

# ─── SHEET 1: Monthly Financial Data ─────────────────────────────────────────
ws1 = wb.active
ws1.title = "Monthly Data"

header_fill  = PatternFill("solid", start_color="1F4E79")
header_font  = Font(bold=True, color="FFFFFF", name="Arial", size=11)
data_font    = Font(name="Arial", size=10)
center       = Alignment(horizontal="center", vertical="center")
currency_fmt = '#,##0'
pct_fmt      = '0.0%'

thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

# Headers
headers = ["Month", "Revenue (₹)", "Expenses (₹)", "Net Profit (₹)", "Profit Margin (%)"]
for col, h in enumerate(headers, 1):
    cell = ws1.cell(row=1, column=col, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center
    cell.border = border

# Monthly data (12 months)
data = [
    ("Jan", 420000, 310000),
    ("Feb", 385000, 295000),
    ("Mar", 510000, 355000),
    ("Apr", 475000, 340000),
    ("May", 540000, 370000),
    ("Jun", 620000, 410000),
    ("Jul", 590000, 395000),
    ("Aug", 650000, 425000),
    ("Sep", 710000, 460000),
    ("Oct", 680000, 440000),
    ("Nov", 760000, 490000),
    ("Dec", 830000, 520000),
]

alt_fill = PatternFill("solid", start_color="AECDEB")

for row_idx, (month, rev, exp) in enumerate(data, 2):
    fill = alt_fill if row_idx % 2 == 0 else PatternFill("solid", start_color="FFFFFF")

    ws1.cell(row=row_idx, column=1, value=month).alignment = center
    ws1.cell(row=row_idx, column=2, value=rev).number_format = currency_fmt
    ws1.cell(row=row_idx, column=3, value=exp).number_format = currency_fmt

    # Net profit formula
    profit_cell = ws1.cell(row=row_idx, column=4)
    profit_cell.value = f"=B{row_idx}-C{row_idx}"
    profit_cell.number_format = currency_fmt

    # Profit margin formula
    margin_cell = ws1.cell(row=row_idx, column=5)
    margin_cell.value = f"=D{row_idx}/B{row_idx}"
    margin_cell.number_format = pct_fmt

    for col in range(1, 6):
        c = ws1.cell(row=row_idx, column=col)
        c.font = data_font
        c.fill = fill
        c.border = border
        c.alignment = center

# Totals row
total_row = 14
ws1.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True, name="Arial", size=10)
ws1.cell(row=total_row, column=1).alignment = center
ws1.cell(row=total_row, column=2, value="=SUM(B2:B13)").number_format = currency_fmt
ws1.cell(row=total_row, column=3, value="=SUM(C2:C13)").number_format = currency_fmt
ws1.cell(row=total_row, column=4, value="=SUM(D2:D13)").number_format = currency_fmt
ws1.cell(row=total_row, column=5, value="=D14/B14").number_format = pct_fmt

total_fill = PatternFill("solid", start_color="D6E4F0")
for col in range(1, 6):
    c = ws1.cell(row=total_row, column=col)
    c.fill = total_fill
    c.font = Font(bold=True, name="Arial", size=10)
    c.border = border
    c.alignment = center

# Column widths
widths = [10, 16, 16, 16, 18]
for col, w in enumerate(widths, 1):
    ws1.column_dimensions[get_column_letter(col)].width = w

ws1.row_dimensions[1].height = 22

# ─── SHEET 2: Department Expenses (for Pie Chart) ─────────────────────────────
ws2 = wb.create_sheet("Dept Expenses")

dept_headers = ["Department", "Annual Budget (₹)", "Actual Spend (₹)", "Variance (₹)"]
for col, h in enumerate(dept_headers, 1):
    cell = ws2.cell(row=1, column=col, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center
    cell.border = border

dept_data = [
    ("Sales & Marketing", 1800000, 1650000),
    ("Operations",        2400000, 2550000),
    ("Human Resources",    900000,  870000),
    ("Technology / IT",  1200000, 1180000),
    ("Finance & Admin",   600000,  580000),
    ("R&D",               750000,  810000),
]

for row_idx, (dept, budget, actual) in enumerate(dept_data, 2):
    fill = alt_fill if row_idx % 2 == 0 else PatternFill("solid", start_color="FFFFFF")
    ws2.cell(row=row_idx, column=1, value=dept).alignment = Alignment(horizontal="left")
    ws2.cell(row=row_idx, column=2, value=budget).number_format = currency_fmt
    ws2.cell(row=row_idx, column=3, value=actual).number_format = currency_fmt
    ws2.cell(row=row_idx, column=4, value=f"=B{row_idx}-C{row_idx}").number_format = currency_fmt
    for col in range(1, 5):
        c = ws2.cell(row=row_idx, column=col)
        c.font = data_font
        c.fill = fill
        c.border = border

dept_widths = [22, 20, 20, 16]
for col, w in enumerate(dept_widths, 1):
    ws2.column_dimensions[get_column_letter(col)].width = w

# ─── SHEET 3: Stock Performance (for Line Chart) ───────────────────────────────
ws3 = wb.create_sheet("Stock Performance")

stock_headers = ["Week", "RELIANCE", "INFY", "TCS", "HDFC"]
for col, h in enumerate(stock_headers, 1):
    cell = ws3.cell(row=1, column=col, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center
    cell.border = border

stock_data = [
    ("W1",  2450, 1520, 3800, 1650),
    ("W2",  2510, 1490, 3850, 1680),
    ("W3",  2480, 1555, 3790, 1700),
    ("W4",  2560, 1610, 3920, 1720),
    ("W5",  2530, 1580, 3870, 1695),
    ("W6",  2620, 1640, 3950, 1760),
    ("W7",  2590, 1700, 4010, 1780),
    ("W8",  2680, 1670, 4080, 1810),
    ("W9",  2750, 1720, 4120, 1840),
    ("W10", 2720, 1760, 4090, 1870),
    ("W11", 2810, 1800, 4200, 1900),
    ("W12", 2870, 1830, 4280, 1940),
]

for row_idx, row_data in enumerate(stock_data, 2):
    fill = alt_fill if row_idx % 2 == 0 else PatternFill("solid", start_color="FFFFFF")
    for col_idx, val in enumerate(row_data, 1):
        c = ws3.cell(row=row_idx, column=col_idx, value=val)
        c.font = data_font
        c.fill = fill
        c.border = border
        c.alignment = center
        if col_idx > 1:
            c.number_format = currency_fmt

stock_widths = [8, 12, 10, 10, 10]
for col, w in enumerate(stock_widths, 1):
    ws3.column_dimensions[get_column_letter(col)].width = w

wb.save("day28_practice_data.xlsx")
print("Practice data file created: day28_practice_data.xlsx")
print("Sheets: Monthly Data | Dept Expenses | Stock Performance")
